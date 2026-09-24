"""
PPTX 解析冒烟测试（自包含：用 python-pptx 构造测试 PPTX）

验证点：
- ZIP 路径：按幻灯片输出 `## 幻灯片 N`、标题 → `###`、正文、GFM 表格；
- 标题文本不重复出现在正文（used 去重）；
- 降级路径（python-pptx）同样可产出内容；
- extract_runs_text 的 <a:br/> 换行与实体反转义。

用法：
    cd backend && PYTHONPATH=. .venv/bin/python scripts/smoke_parse_pptx.py
"""
import io

from pptx import Presentation
from pptx.util import Inches

from app.utils.parsers.pptx import (
    _parse_pptx_with_python_pptx,
    decode_xml,
    extract_runs_text,
    parse_pptx,
)


def build_test_pptx() -> bytes:
    prs = Presentation()

    # 第 1 页：标题 + 正文
    slide1 = prs.slides.add_slide(prs.slide_layouts[1]) # Title and Content
    slide1.shapes.title.text = "产品介绍"
    slide1.placeholders[1].text = "这是第一页的正文内容。"

    # 第 2 页：标题 + 正文 + 表格
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "数据对比"
    slide2.placeholders[1].text = "以下是两组数据。"
    table_shape = slide2.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(4), Inches(1.2)
    )
    table = table_shape.table
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "100"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def main() -> None:
    pptx_bytes = build_test_pptx()

    # 1) 主路径（ZIP/XML）
    md = parse_pptx(pptx_bytes)
    print("--- 主路径解析结果 ---")
    print(md)
    print("----------------------")
    assert "## 幻灯片 1" in md and "## 幻灯片 2" in md, "幻灯片章节丢失"
    assert "### 产品介绍" in md, "标题占位符未转 ###"
    assert "第一页的正文内容" in md, "正文丢失"
    assert "### 数据对比" in md
    assert "以下是两组数据。" in md
    assert "| 指标 | 数值 |" in md and "| --- | --- |" in md, "表格未转 GFM"
    assert "| A | 100 |" in md
    # 标题文本不应重复出现在正文（used 去重）
    assert md.count("产品介绍") == 1, f"标题重复: {md}"

    # 2) 降级路径（python-pptx）直接调用
    fallback = _parse_pptx_with_python_pptx(pptx_bytes)
    assert "## 幻灯片 1" in fallback and "产品介绍" in fallback, fallback
    print("2) 降级路径 OK")

    # 3) run 提取：<a:br/> → 换行
    runs = extract_runs_text(
        "<a:p><a:r><a:t>你好</a:t></a:r><a:r><a:br/></a:r>"
        "<a:r><a:t>世界</a:t></a:r></a:p>"
    )
    assert runs == "你好\n世界", repr(runs)
    # 4) 实体反转义
    assert decode_xml("a &lt; b &amp; c &quot;d&quot;") == 'a < b & c "d"'
    print("3/4) runs 提取与实体反转义 OK")

    print("\nPPTX 解析冒烟测试全部通过 ✅")


if __name__ == "__main__":
    main()
